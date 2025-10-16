import io
from typing import Optional
from fastapi import UploadFile


async def _read_all(file: UploadFile) -> bytes:
    data = await file.read()
    await file.seek(0)
    return data


async def extract_text_from_file(file: UploadFile) -> str:
    name = (file.filename or "").lower()
    content_type = file.content_type or ""
    data = await _read_all(file)

    # Simple routing by extension/content-type
    if name.endswith(".pdf") or content_type == "application/pdf":
        return extract_text_from_pdf_bytes(data)
    if name.endswith(".docx"):
        return extract_text_from_docx_bytes(data)
    if name.endswith(".pptx"):
        return extract_text_from_pptx_bytes(data)
    if name.endswith(".txt"):
        return data.decode(errors="ignore")
    if name.endswith(".csv") or content_type in ("text/csv",):
        return data.decode(errors="ignore")
    if name.endswith(".xlsx"):
        return extract_text_from_xlsx_bytes(data)

    # Fallback to bytes decode (may be garbage for binaries)
    try:
        return data.decode(errors="ignore")
    except Exception:
        return ""


def extract_text_from_pdf_bytes(data: bytes) -> str:
    try:
        import pdfminer.high_level as pdfminer
        text = pdfminer.extract_text(io.BytesIO(data))
        return text or ""
    except Exception:
        # fallback: pypdf
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(data))
            buf = []
            for page in reader.pages:
                buf.append(page.extract_text() or "")
            return "\n".join(buf)
        except Exception:
            return ""


def extract_text_from_docx_bytes(data: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        return ""


def extract_text_from_pptx_bytes(data: bytes) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texts.append(shape.text)
        return "\n".join(texts)
    except Exception:
        return ""


def extract_text_from_xlsx_bytes(data: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                lines.append("\t".join(vals))
        return "\n".join(lines)
    except Exception:
        return ""
